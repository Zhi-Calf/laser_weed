"""Build the operator guide and dimensioned calibration print assets.
Run with the bundled document Python; screenshot assets are captured separately.
"""
from pathlib import Path
import math
import os
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pdf2image import convert_from_path

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/calibration-guide-v4'; AS=OUT/'assets'; AS.mkdir(parents=True,exist_ok=True)
pdfmetrics.registerFont(TTFont('GuideCJK', os.environ.get('CALIBRATION_GUIDE_FONT', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf')))
INK='#18322d'; GREEN='#26744b'; ORANGE='#a95013'
def txt(c,x,y,s,size=14,color=INK):
 c.setFillColor(color);c.setFont('GuideCJK',size);c.drawString(x,y,s)
def line(c,x,y,X,Y,color=INK,width=1.6):
 c.setStrokeColor(color);c.setLineWidth(width);c.line(x,y,X,Y)
def arrow(c,x,y,X,Y,color=INK):
 line(c,x,y,X,Y,color,2);a=math.atan2(Y-y,X-x)
 for d in [-.5,.5]:line(c,X,Y,X-10*math.cos(a+d),Y-10*math.sin(a+d),color,2)
def rect(c,x,y,w,h,fill='#f0f4f0'):
 c.setFillColor(fill);c.setStrokeColor(INK);c.setLineWidth(1);c.rect(x,y,w,h,fill=1)
def target(c,x,y,n,check=False,size=20):
 c.setFillColor('white');c.setStrokeColor('black');c.rect(x-size/2,y-size/2,size,size,fill=1)
 c.setFillColor('black');c.rect(x-size/2,y,size/2,size/2,fill=1,stroke=0);c.rect(x,y-size/2,size/2,size/2,fill=1,stroke=0)
 txt(c,x+size/2+5,y-3,str(n),14,ORANGE if check else GREEN)
def title(c,s,sub):
 txt(c,35,555,s,24,'black');txt(c,35,525,sub,13)
def footer(c,n):txt(c,35,22,'操作示意图  不按比例绘制  不能用作内参打印板',11);txt(c,760,22,str(n),11)

c=canvas.Canvas(str(OUT/'标定操作示意图_v4.pdf'),pagesize=landscape(A4))
title(c,'内参标定 正视图','相机保持固定；九宫格表示板在画面中的位置，不是在地面布点。')
for row in range(3):
 for col in range(3):
  x=65+col*170;y=125+(2-row)*115
  rect(c,x,y,160,105,'#f5f7f5')
  txt(c,x+8,y+82,['左上','上中','右上','左中','中心','右中','左下','下中','右下'][row*3+col],12)
  c.saveState();c.translate(x+80,y+45);c.rotate([-12,0,12][col]);rect(c,-19,-27,38,54,'white')
  for rr in range(5):
   for cc in range(3):
    if (rr+cc)%2==0:c.setFillColor('black');c.rect(-15+cc*10,-23+rr*9,10,9,fill=1,stroke=0)
  c.restoreState()
txt(c,610,420,'每个位置',18);txt(c,610,385,'1 摆好并保持不动');txt(c,610,350,'2 点击拍照');txt(c,610,315,'3 看回显是否清晰');txt(c,610,280,'4 满意后点下一位置');txt(c,610,225,'近 中 远各一轮');txt(c,610,190,'穿插向左右上下倾斜');txt(c,65,75,'板应完整入镜；不要挡住图案。示意方块不是实际 ChArUco 图案。',14);footer(c,1);c.showPage()

title(c,'内参标定 俯视图','改变板与相机的距离和左右位置；举板高度、上下倾斜参照正视图。')
rect(c,365,65,95,45);txt(c,382,80,'相机固定',13)
for x in [80,745]:line(c,412,110,x,455,'#b3c8bd')
for y,lab,xs in [(205,'近',[330,410,490]),(315,'中',[235,410,585]),(425,'远',[150,410,670])]:
 for x in xs:
  c.saveState();c.translate(x,y);c.rotate(15 if x<410 else -15 if x>410 else 0);rect(c,-28,-7,56,14);c.restoreState()
 txt(c,50,y-4,lab,19)
arrow(c,780,195,780,440);txt(c,690,150,'距离逐步增加',14)
txt(c,85,485,'同一块板依次移动，不是同时摆九块板。',16)
txt(c,70,40,'距离不必输入页面；以图案清晰、角点可辨、画面中大小不同为准。',14);footer(c,2);c.showPage()

title(c,'BEV 标定 地面俯视图','一次布置全部标记；以后轮轴中心的地面投影为原点。')
rect(c,260,65,110,120);txt(c,280,150,'车辆',15)
line(c,245,88,385,88,width=4);txt(c,390,82,'后轴',12)
arrow(c,315,88,315,490);arrow(c,315,88,95,88)
txt(c,325,475,'前 +x',14);txt(c,90,65,'左 +y',14);txt(c,325,102,'O',16)
for row,y in enumerate([225,295,365,435]):
 target(c,145,y,2*row+1);target(c,475,y,2*row+2)
for y,n in [(260,9),(400,11)]:target(c,230,y,n,True);target(c,400,y,n+1,True)
txt(c,550,465,'1–8 标定点',17,GREEN);txt(c,550,432,'9–12 检查点',17,ORANGE)
txt(c,550,395,'所有点必须同时入镜');txt(c,550,365,'覆盖实际作业近远范围');txt(c,550,335,'位置按现场视野调整')
# measurement inset
arrow(c,670,145,670,280);arrow(c,670,145,560,145)
line(c,670,245,575,245,'#a95013');line(c,575,245,575,145,'#aaaaaa')
c.setFillColor(ORANGE);c.circle(575,245,4,fill=1);txt(c,550,263,'P',14);txt(c,681,238,'Q',14);txt(c,679,143,'O',14)
txt(c,687,195,'x',14);txt(c,610,253,'y',14);txt(c,550,115,'O 到 Q 量前向 x');txt(c,550,87,'Q 到 P 量横向 y')
txt(c,65,40,'左侧 y 为正，右侧 y 为负；两个方向互相垂直，不量 O 到 P 的斜距。',13);footer(c,3);c.showPage()

title(c,'BEV 标定 相机画面示意','页面照片已去畸变；先选表格点号，再点击同号棋盘的中心交汇点。')
rect(c,55,100,540,390,'#ededdf')
for yy,spread,sz,row in [(150,210,38,0),(235,160,30,1),(320,115,24,2),(405,80,18,3)]:
 target(c,325-spread,yy,row*2+1,size=sz);target(c,325+spread,yy,row*2+2,size=sz)
for yy,spread,n,sz in [(195,85,9,28),(365,50,11,20)]:target(c,325-spread,yy,n,True,sz);target(c,325+spread,yy,n+1,True,sz)
txt(c,70,465,'远',15);txt(c,70,115,'近',15)
target(c,705,340,'',size=110);arrow(c,780,245,707,337,ORANGE)
txt(c,620,215,'点击黑白交汇中心',16)
txt(c,620,180,'不是纸张中心或编号',13)
txt(c,620,140,'放大后仍须能辨认中心',13)
txt(c,65,60,'此图仅说明对应关系。实际画面形状由相机位置决定；不要照抄示意图坐标。',13);footer(c,4);c.save()

# A4 dimensioned print markers, one per sheet.
c=canvas.Canvas(str(OUT/'BEV地面棋盘标记_A4_01至12.pdf'),pagesize=A4)
for n in range(1,13):
 txt(c,25*mm,268*mm,f'{n:02d}  '+('标定点' if n<=8 else '检查点'),30,'black')
 txt(c,25*mm,250*mm,'BEV 地面标记  不是内参标定板',14,'black')
 x=25*mm;y=70*mm;s=160*mm
 c.setStrokeColor('black');c.setFillColor('white');c.setLineWidth(.5);c.rect(x,y,s,s,fill=1)
 c.setFillColor('black');c.rect(x,y+s/2,s/2,s/2,fill=1,stroke=0);c.rect(x+s/2,y,s/2,s/2,fill=1,stroke=0)
 txt(c,25*mm,55*mm,'测量与点选位置：四个方格交汇中心',13,'black')
 txt(c,25*mm,45*mm,'A4 单面打印，实际大小 100%，不要适应页面。',11,'black')
 txt(c,25*mm,37*mm,'保持平整并固定四角；不要在交汇中心贴胶带。',11,'black')
 line(c,25*mm,25*mm,125*mm,25*mm,'black',1)
 for xx in [25,125]:line(c,xx*mm,23*mm,xx*mm,27*mm,'black',1)
 txt(c,130*mm,24*mm,'100 mm 校验线',11,'black')
 c.showPage()
c.save()
# Raster diagrams for insertion, not substitutes for printable vector targets.
pdf=convert_from_path(OUT/'标定操作示意图_v4.pdf', dpi=144, poppler_path=os.environ.get('POPPLER_PATH'))
for i,page in enumerate(pdf):page.save(AS/f'diagram-{i+1}.png')

# Word manual. One task per page, actual current UI screenshots.
im=Image.open(AS/'capture.png'); combo=Image.new('RGB',(im.width*2+20,735),'white');combo.paste(im.crop((0,0,im.width,735)),(0,0));combo.paste(im.crop((0,735,im.width,im.height)),(im.width+20,0));combo.save(AS/'capture-pair.png')
doc=Document();sec=doc.sections[0]
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.5);sec.left_margin=sec.right_margin=Cm(1.7)
for name,size in [('Normal',11),('Title',23),('Subtitle',12),('Heading 1',18),('Heading 2',12),('Caption',9)]:
 s=doc.styles[name];s.font.name='Arial Unicode MS';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
 s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.1
for name in ['Heading 1','Heading 2']:doc.styles[name].paragraph_format.space_before=Pt(5)
foot=sec.footer.paragraphs[0];foot.alignment=2
foot.add_run('田间标定操作手册 v4.0  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(s,style=None):return doc.add_paragraph(s,style)
def h(s):doc.add_heading(s,2)
def page(s):doc.add_page_break();doc.add_heading(s,1)
def pic(name,width=6.6):
 par=doc.add_paragraph();par.alignment=1;par.add_run().add_picture(str(AS/name),width=Inches(width))
def caption(s):p(s,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for cells,vals in [(t.rows[0].cells,headers)]+[(t.add_row().cells,row) for row in rows]:
  for i,(cell,val) in enumerate(zip(cells,vals)):
   cell.text=str(val);cell.vertical_alignment=1
   if widths:cell.width=Cm(widths[i])
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   tc=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF0' if ri==0 else 'FFFFFF');tc.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
   tc.append(margins)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(3)
    for run in par.runs:run.font.size=Pt(10);run.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge)
  for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
  borders.append(el)
 t._tbl.tblPr.append(borders)
 return t

p('田间相机标定操作手册','Title');p('内参拍板与装车地面 BEV 标定','Subtitle');p('版本 4.0   2026年10月29日')
p('本手册供现场同事使用。技术人员先准备好可访问的田间标定台，操作人员按页面拍照、测量、填写和保存即可，不需要运行代码。')
p('本版采用已实现的流程：内参使用 ChArUco 板；BEV 使用同一张照片中的多个实测地面标记。请使用本版配套图，不再使用旧版的多位置移动标定板外参图。')
h('开始前准备')
p('已打印且贴在硬板上的 ChArUco 图；A4 地面棋盘标记 01–12；卷尺、直角尺、细绳或胶带、铅垂工具；可访问页面的电脑。地面棋盘标记平整固定即可，不必抬高；内参板与地面标记不能混用。')
h('打开页面')
p('Windows 或 Ubuntu 笔记本按部署人员交接的方式连接设备，在浏览器打开下方地址。有配套显示器时，也使用已经配置好的浏览器入口。网口、地址和服务启动由部署人员确认，不自行猜测。')
p('页面地址：________________________________________________')
p('联系人：________________    日期及车辆编号：________________')
h('两项任务按顺序做')
table(['任务','你需要做什么','完成后得到什么'],[['内参','手持同一块板，变化位置、距离和倾斜角度','用于去畸变的内参文件'],['BEV','相机最终装车，测量地面标记并点选照片','以后轴地面原点为参考的映射']],[2,9,6.4])
p('车辆静止、相机固定后再进行 BEV 标定。内参通过复核后才能用于 BEV；相机焦距、分辨率或裁剪改变时，应先联系负责人。')
p('本手册页面截图来自当前程序的模拟模式，只说明操作位置。图中数值和检查结果不能作为现场精度证明；模拟模式不能设为正式算法参数。','Caption')

page('一 建立现场会话')
p('进入田间标定台，点击“刷新状态”。现场应使用真机模式；如果显示模拟模式，请联系部署人员。首页预览可能是文件预览，真正拍下的照片以拍照后的回显为准。')
pic('session.png',4.8)
caption('当前页面截图。尺寸单位为毫米；示例值来自原板历史记录，新板必须实测。')
p('① 填操作人员。② 量连续 5 格的总长度再除以 5，分别填写横向、纵向方格边长，不包含外白边。③ 填小编码标记的外边长与 100 mm 校验线实际长度。④ 确认板平整，点击“开始新会话”。')
p('原板历史尺寸为 56.5 mm、56.5 mm、42.2 mm、100 mm；仅在确认同一块板且尺寸未变时使用。开始会话不等于开始拍照。一次会话由一位操作人员操作。')

page('二 内参板如何摆放')
pic('diagram-1.png',6.0)
pic('diagram-2.png',6.0)
p('近、中、远是相对距离，不需要输入米数。每轮覆盖九宫格，并穿插左右、上下倾斜。板保持清晰、完整入镜，避免手挡图案、反光和运动模糊。')

page('三 内参拍照与重试')
pic('capture-pair.png',6.45)
caption('当前页面截图分为左右两栏：左为采集位置，右为回显与按钮。模拟图案不能用于真实内参。')
p('① 在下拉框选“拟合照片”，按页面提示摆板。② 点击“准备好了，拍照”。③ 查看照片，不满意点“不满意，重拍这张”；满意点“照片满意，下一个位置”。④ 依次完成近、中、远共 27 张。')
p('⑤ 下拉框切换到“独立验证”，重新摆放拍 9 张，不要复制拟合照片。⑥ 点击“使用本次照片计算内参”，等待候选结果。拟合和验证是两组照片，这与 BEV 同一张照片内的检查点不同。')
h('计算失败后')
p('留在原会话，按提示补拍或重拍，再点击计算；照片不会被清空。九宫格可选择当前轮的位置；需要回到后面的轮次时，从前面的槽位逐步拍照并点“下一位置”，数量不会因重拍重复增加。页面目前没有跨轮次一键跳转。')
p('计算成功后本会话的内参采集锁定。需要重新做成功的标定时，新建会话。候选内参需负责人复核清晰度、边缘去畸变效果和误差，不能只看“计算成功”。')

page('四 建立后轴测量基准')
p('相机安装在最终位置，车辆停在平整硬地面。找出左右后轮轮心在地面的垂直投影，两点中点作为 O；以车辆纵向中心线确定向前方向，不用转向轮当前指向代替车体方向。')
pic('diagram-3.png',6.55)
p('沿前向中心线从 O 量到垂足 Q，得到 x；从 Q 垂直量到标记中心 P，得到 y。面向车头，左侧填正值，右侧填负值。用直角尺确定横向，不量 O 到 P 的斜距。')
table(['标记实际位置','前向 x / m','左向 y / m'],[['后轴前方 2 m 左侧 0.6 m','2.000','0.600'],['后轴前方 3 m 右侧 0.8 m','3.000','-0.800']],[10,3.7,3.7])
p('上表仅解释填写方式，不是强制摆放距离。原点或中心线若测错，整套结果都会偏移，不能靠低拟合误差发现所有基准错误。')

page('五 摆放地面棋盘并测量')
p('打印配套 A4 文件全部 12 页，建议单面、100% 实际大小，不选择“适应页面”。每页一个 160 mm 棋盘，单格 80 mm。校验线应为 100 mm；BEV 使用实测中心坐标，不用格子边长推算车体距离。')
p('编号 01–08 用作标定点，09–12 用作检查点；全部一次布置、同一张照片拍下，不另安排检查拍摄。标记应覆盖实际作业范围，保持所有中心清晰可见。远处辨不清时放大打印标记或调整有效测量范围，不凭感觉点选。')
pic('diagram-4.png',6.55)
p('把四角固定，避免起翘、折弯、滑动和反光。编号朝向方便辨认即可；不要在黑白交汇中心贴胶带。记录每个中心的 x、y，使用本手册最后一页记录表。')
p('本版不移动一块大板分次拍摄，也没有自动识别标记功能。每个标记必须在页面手动点选。')

page('六 拍摄地面照片')
pic('bev-prepare.png',6.55)
caption('当前页面截图。下拉框默认可能是预配置内参，使用本次结果时需主动切换。')
p('① 找到“把画面中的地面点，对齐后轮轴中心”。② 在“使用哪份内参”选择“本次会话计算的内参”；若负责人已配置并复核另一份内参，可选预配置项。③ 确认后轴原点、地面平整、相机固定，以及内参已复核，勾选两个确认框。')
p('④ 点击“准备好了，拍摄地面照片”，等待去畸变后的冻结照片。确认所有标记中心清楚、没有被车身遮挡；这张照片才是后续点选依据。')
h('是否需要重新拍照')
p('只填错坐标或点错像素时，不必重拍，直接改点。标记移动、照片模糊或相机移动时才重新拍摄；相机位置改变后应重新核对整套 BEV。重拍成功会清空当前点位与结果，请先导出需保留的记录。重拍处理失败时上一份照片与点位仍保留。')

page('七 输入坐标并点选')
pic('bev-points.png',6.55)
caption('当前页面截图。棋盘图像为模拟数据，现场点击纸质棋盘交汇中心；不要照抄图中坐标。')
p('① 点击表格第 1 点，确认用途为“标定”。② 填记录表对应的前向 x、左向 y，单位米。③ 在照片上点击 01 号棋盘黑白交汇中心。④ 对 02–12 号依次操作，09–12 的用途保持“检查”。')
p('点选错误时重新点击正确中心；坐标错误时直接改输入框。不要把同一个位置同时用作标定和检查点。全部完成后先点“保存点位”，再点“计算 BEV 并检查误差”。')
p('计算前程序至少要求 6 个标定点、4 个检查点；本手册建议 8+4。所有点不要排列在同一直线上。填入超范围、重复或缺失数据时，按提示修改后重试，不需要新建会话。')

page('八 查看结果并设为算法参数')
pic('bev-result.png',6.2)
caption('当前程序模拟结果截图。图中的极小误差来自测试数据，不代表现场精度；模拟模式保存按钮不可用。')
p('查看鸟瞰图、编号与各点误差。画面上方应为车头方向、左侧为车辆左侧。展开“查看每个点的误差”，先排查编号、左右符号、测量基准和点选是否正确。')
p('当前几何门槛：拟合内点比例至少 75%，检查点 RMSE 不超过 5 cm、最大误差不超过 10 cm。这只是程序门槛，不是导航精度承诺。未通过时先重测或重选异常点，不为通过而随意删除检查点。')
p('真实数据通过检查，且负责人确认内参、后轴参考点与已有轴距补偿一致后，勾选人工复核，点击“复核后设为算法参数”。软件保存新版本并保留旧版本；候选文件不会直接当成正式参数。')
p('保存不会自动重启车辆程序。请把保存成功信息交给集成人员，由其确认算法重新初始化并加载新版本。')

page('九 失败处理与结果交接')
table(['遇到的情况','怎么处理'],[['内参计算失败','留在原会话补拍或重拍，再点计算；失败记录会保留。'],['BEV 计算失败','点位保留，旧结果隐藏。修改坐标或点选后重算，无需新建会话。'],['提示无法刷新最新状态','先恢复网络。刷新可能丢失未保存的本地修改，先记下；恢复后核对照片和点位再算。'],['拍照失败','检查相机连接或是否被其他程序占用，联系部署人员后重试。'],['设为算法参数按钮不可点','检查是否真机模式、是否通过几何检查、是否有未计算的修改、是否已勾选复核。'],['保存失败','已有正式版本仍保留。检查连接或联系部署人员；重新核对并勾选复核后再保存。']],[5.1,12.3])
pic('export.png',5.3)
caption('当前页面的会话与导出入口。')
p('点击“导出测试包”，再点“下载 ZIP”。交接 ZIP、地面测量记录表、车辆与相机安装照片、保存成功截图及会话编号。错误日志已包含在包内，现场人员不必手工查文件。')
p('只有看到加载确认并完成实车距离与方向检查，才算部署验证完成。单纯算出参数、下载文件或通过模拟测试均不等于实车验收。')

page('十 地面标记测量记录表')
p('车辆编号：________________   操作人员：________________')
p('日期：____________________   会话编号：________________')
p('原点 O 已核对 □    前向中心线已核对 □    所有标记中心清晰入镜 □')
table(['点号','用途','前向 x / m','左向 y / m','照片对应已核对'],[[f'{i:02d}','标定' if i<=8 else '检查','________','________','□'] for i in range(1,13)],[1.5,2.2,4.4,4.4,4.9])
p('填写实测值，不照抄示例。左正右负；不量斜距；纸面编号必须与页面点号一致。')
p('检查点 RMSE：________ cm    最大误差：________ cm')
p('已保存正式参数 □    已导出 ZIP □    已交接加载确认 □')
p('备注：____________________________________________________')
p('__________________________________________________________')

doc.save(OUT/'田间标定台现场操作手册_v4.0.docx')
print('Created manual, diagrams and 12-page printable markers in',OUT)
